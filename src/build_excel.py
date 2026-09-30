"""Builds outputs/risk_dashboard.xlsx.

Data tabs hold values pulled from the database / model outputs. Everything on the
Risk and Summary tabs is a live formula, so changing weights on the Inputs tab updates it all.
"""
import sqlite3
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import DB_PATH, OUT, PORTFOLIO, UNIVERSE

N_DAYS = 504
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
WARN = Font(name=FONT, bold=True, color="C00000")
thin = Side(style="thin", color="BFBFBF")

PCT = "0.00%;(0.00%);-"
PCT1 = "0.0%;(0.0%);-"
USD = "$#,##0;($#,##0);-"


def style_range(ws, ref, font=None, fmt=None, fill=None, align=None):
    for row in ws[ref]:
        for c in row:
            if font: c.font = font
            if fmt: c.number_format = fmt
            if fill: c.fill = fill
            if align: c.alignment = align


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_df(ws, df, top, left=1, fmts=None):
    header(ws, top, list(df.columns), left)
    for i, row in enumerate(df.itertuples(index=False), start=top + 1):
        for j, v in enumerate(row):
            c = ws.cell(row=i, column=left + j, value=v)
            c.font = BLUE
            if fmts and df.columns[j] in fmts:
                c.number_format = fmts[df.columns[j]]
    return top + len(df)


def load_returns():
    con = sqlite3.connect(DB_PATH)
    tickers = PORTFOLIO + ["SPY"]
    q = f"SELECT ticker, date, ret FROM daily_returns WHERE ticker IN ({','.join('?' * len(tickers))})"
    df = pd.read_sql(q, con, params=tickers)
    source = con.execute("SELECT value FROM meta WHERE key = 'source'").fetchone()[0]
    con.close()
    wide = df.pivot(index="date", columns="ticker", values="ret")[tickers].dropna().tail(N_DAYS)
    return wide, source


def main():
    rets, source = load_returns()
    synthetic = source == "synthetic"
    n = len(rets)
    first, last = 6, 6 + n - 1                       # data rows on the Returns sheet
    T = PORTFOLIO
    k = len(T)
    stock_cols = [L(2 + i) for i in range(k)]        # B..K
    spy_col, port_col = L(2 + k), L(3 + k)           # L, M
    growth_col, max_col, dd_col = L(4 + k), L(5 + k), L(6 + k)   # N, O, P
    rng = lambda col: f"Returns!${col}${first}:${col}${last}"

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_in = wb.create_sheet("Inputs")
    ws_ret = wb.create_sheet("Returns")
    ws_risk = wb.create_sheet("Risk")
    ws_sec = wb.create_sheet("Sector_Stats")
    ws_reg = wb.create_sheet("Regime_Analysis")
    ws_mod = wb.create_sheet("Model_Results")

    banner = ("SYNTHETIC DEMO DATA - rerun the pipeline with real data (see README) before using any numbers"
              if synthetic else f"Data source: {source}")

    # ---------------- Inputs ----------------
    ws = ws_in
    ws["A1"], ws["A1"].font = "Portfolio inputs", TITLE
    ws["A2"], ws["A2"].font = banner, WARN if synthetic else BLACK
    labels = [("Portfolio value ($)", 1_000_000, USD, "Assumption: notional size for $ risk numbers"),
              ("VaR confidence level", 0.95, "0.0%", "Assumption: 95% is a common desk convention"),
              ("Risk-free rate (annual)", 0.03, "0.00%", "Assumption: flat rate used in Sharpe ratio"),
              ("Trading days per year", 252, "0", "Convention"),
              ("VaR horizon (days)", 10, "0", "Used for the square-root-of-time VaR")]
    for i, (lab, val, fmt, note) in enumerate(labels, start=3):
        ws.cell(row=i, column=1, value=lab).font = BLACK
        c = ws.cell(row=i, column=2, value=val)
        c.font, c.fill, c.number_format = BLUE, INPUT_FILL, fmt
        ws.cell(row=i, column=3, value=note).font = Font(name=FONT, italic=True, color="595959")
    header(ws, 9, ["Ticker", "Name", "Sector", "Weight"])
    for i, t in enumerate(T):
        r = 10 + i
        name, sector = UNIVERSE[t]
        ws.cell(row=r, column=1, value=t).font = BLUE
        ws.cell(row=r, column=2, value=name).font = BLUE
        ws.cell(row=r, column=3, value=sector).font = BLUE
        c = ws.cell(row=r, column=4, value=1 / k)
        c.font, c.fill, c.number_format = BLUE, INPUT_FILL, "0.0%"
    w_first, w_last = 10, 9 + k
    ws.cell(row=w_last + 1, column=3, value="Total").font = BOLD
    c = ws.cell(row=w_last + 1, column=4, value=f"=SUM(D{w_first}:D{w_last})")
    c.font, c.number_format = BOLD, "0.0%"
    ws.cell(row=w_last + 1, column=5,
            value=f'=IF(ABS(D{w_last + 1}-1)<0.0001,"OK","Weights must sum to 100%")').font = BOLD
    ws.cell(row=w_last + 3, column=1, value="Legend").font = BOLD
    ws.cell(row=w_last + 4, column=1, value="Blue text = hardcoded input or data").font = BLUE
    ws.cell(row=w_last + 5, column=1, value="Black text = formula").font = BLACK
    ws.cell(row=w_last + 6, column=1, value="Green text = link from another sheet").font = GREEN
    c = ws.cell(row=w_last + 7, column=1, value="Yellow fill = cells to change (weights, portfolio value, etc.)")
    c.font, c.fill = BLUE, INPUT_FILL
    ws.cell(row=w_last + 8, column=1,
            value="Tip: try Data > Solver to maximise the Sharpe ratio on Risk!B9 by changing the weights.").font = BLACK
    for col, w in zip("ABCDE", [26, 22, 26, 12, 28]):
        ws.column_dimensions[col].width = w

    # ---------------- Returns ----------------
    ws = ws_ret
    ws["A1"], ws["A1"].font = f"Daily simple returns, last {n} trading days", TITLE
    ws["A2"], ws["A2"].font = "Source: daily_returns view in the SQL database (adjusted close). Data in blue.", BLACK
    ws["A3"], ws["A3"].font = "Weight", BOLD
    for i in range(k):
        c = ws[f"{stock_cols[i]}3"]
        c.value, c.font, c.number_format = f"=Inputs!D{w_first + i}", GREEN, "0.0%"
    header(ws, 5, ["Date"] + [None] * k + ["SPY", "Portfolio", "Growth of 1", "Running max", "Drawdown"])
    for i in range(k):
        c = ws[f"{stock_cols[i]}5"]
        c.value = f"=Inputs!A{w_first + i}"
    for i, (d, row) in enumerate(rets.iterrows()):
        r = first + i
        ws.cell(row=r, column=1, value=d).font = BLUE
        for j, t in enumerate(T + ["SPY"]):
            c = ws.cell(row=r, column=2 + j, value=float(row[t]))
            c.font, c.number_format = BLUE, "0.00%"
        ws[f"{port_col}{r}"] = f"=SUMPRODUCT($B$3:${stock_cols[-1]}$3,B{r}:{stock_cols[-1]}{r})"
        ws[f"{growth_col}{r}"] = f"=1+{port_col}{r}" if i == 0 else f"={growth_col}{r - 1}*(1+{port_col}{r})"
        ws[f"{max_col}{r}"] = f"=MAX(1,{growth_col}{r})" if i == 0 else f"=MAX({max_col}{r - 1},{growth_col}{r})"
        ws[f"{dd_col}{r}"] = f"={growth_col}{r}/{max_col}{r}-1"
        for col, fmt in [(port_col, "0.00%"), (growth_col, "0.000"), (max_col, "0.000"), (dd_col, "0.00%")]:
            ws[f"{col}{r}"].font, ws[f"{col}{r}"].number_format = BLACK, fmt
    style_range(ws, f"A{first}:A{last}", fmt="yyyy-mm-dd")
    ws.freeze_panes = f"B{first}"
    ws.column_dimensions["A"].width = 12
    for j in range(2, 7 + k):
        ws.column_dimensions[L(j)].width = 11

    # ---------------- Risk ----------------
    ws = ws_risk
    ws["A1"], ws["A1"].font = "Portfolio risk", TITLE
    ws["A2"], ws["A2"].font = banner, WARN if synthetic else BLACK
    header(ws, 4, ["Metric", "Value", "Note"])
    P = rng(port_col)
    conf, ndays, rf, val, hz = "Inputs!$B$4", "Inputs!$B$6", "Inputs!$B$5", "Inputs!$B$3", "Inputs!$B$7"
    metrics = [
        ("Observations", f"=COUNT({P})", "0", "Trading days in the sample"),
        ("Mean daily return", f"=AVERAGE({P})", "0.000%", ""),
        ("Annualised return", f"=B6*{ndays}", PCT, "Arithmetic: mean daily x days per year"),
        ("Annualised volatility", f"=STDEV({P})*SQRT({ndays})", PCT, "Sample st. dev. x sqrt(days)"),
        ("Sharpe ratio", f"=(B7-{rf})/B8", "0.00", "Uses the risk-free rate on Inputs"),
        ("Best day", f"=MAX({P})", PCT, ""),
        ("Worst day", f"=MIN({P})", PCT, ""),
        ("Maximum drawdown", f"=MIN({rng(dd_col)})", PCT, "Peak to trough, from Returns tab"),
        ("Beta vs SPY", f"=SLOPE({P},{rng(spy_col)})", "0.00", "Regression slope on daily returns"),
        ("1-day historical VaR", f"=-PERCENTILE({P},1-{conf})", PCT, "Loss not exceeded at the confidence level"),
        ("1-day parametric VaR", f"=-(B6+NORMSINV(1-{conf})*STDEV({P}))", PCT, "Assumes normal returns"),
        ("1-day CVaR (expected shortfall)", f'=-AVERAGEIF({P},"<="&-B14)', PCT, "Average loss on days beyond the historical VaR"),
        ("1-day historical VaR ($)", f"=B14*{val}", USD, ""),
        ("1-day CVaR ($)", f"=B16*{val}", USD, ""),
        ("Horizon parametric VaR ($)", f"=B15*SQRT({hz})*{val}", USD, "Square-root-of-time, assumes i.i.d. returns"),
        ("Days worse than hist. VaR", f'=COUNTIF({P},"<"&-B14)', "0", "Should be roughly (1 - confidence) x observations"),
    ]
    for i, (lab, f, fmt, note) in enumerate(metrics, start=5):
        ws.cell(row=i, column=1, value=lab).font = BLACK
        c = ws.cell(row=i, column=2, value=f)
        c.font, c.number_format = BLACK, fmt
        ws.cell(row=i, column=3, value=note).font = Font(name=FONT, italic=True, color="595959")

    # per stock table
    s_top = 24
    ws.cell(row=s_top - 1, column=1, value="Per-stock statistics").font = BOLD
    header(ws, s_top, ["Ticker", "Sector", "Weight", "Ann. return", "Ann. volatility",
                       "Beta (Excel, sample)", "Beta (SQL, full history)", "Correlation to SPY"])
    sql_beta = pd.read_csv(OUT / "sql" / "q5_beta_vs_spy.csv").set_index("ticker")["beta"]
    for i, t in enumerate(T):
        r = s_top + 1 + i
        col = stock_cols[i]
        ws[f"A{r}"], ws[f"A{r}"].font = f"=Inputs!A{w_first + i}", GREEN
        ws[f"B{r}"], ws[f"B{r}"].font = f"=Inputs!C{w_first + i}", GREEN
        ws[f"C{r}"], ws[f"C{r}"].font = f"=Inputs!D{w_first + i}", GREEN
        ws[f"D{r}"] = f"=AVERAGE({rng(col)})*{ndays}"
        ws[f"E{r}"] = f"=STDEV({rng(col)})*SQRT({ndays})"
        ws[f"F{r}"] = f"=SLOPE({rng(col)},{rng(spy_col)})"
        ws[f"G{r}"] = float(sql_beta[t])
        ws[f"H{r}"] = f"=CORREL({rng(col)},{rng(spy_col)})"
        for c_, fmt in [("C", PCT1), ("D", PCT1), ("E", PCT1), ("F", "0.00"), ("G", "0.00"), ("H", "0.00")]:
            if c_ != "C":
                ws[f"{c_}{r}"].font = BLUE if c_ == "G" else BLACK
            ws[f"{c_}{r}"].number_format = fmt
    s_first, s_last = s_top + 1, s_top + k
    r = s_last + 1
    ws[f"A{r}"], ws[f"A{r}"].font = "Weighted beta check (should equal B13)", BOLD
    ws[f"F{r}"] = f"=SUMPRODUCT(C{s_first}:C{s_last},F{s_first}:F{s_last})"
    ws[f"F{r}"].font, ws[f"F{r}"].number_format = BLACK, "0.00"

    # correlation matrix
    c_top = s_last + 5
    ws.cell(row=c_top - 1, column=1, value="Correlation matrix (daily returns)").font = BOLD
    header(ws, c_top, [None] + [None] * k)
    for j in range(k):
        ws.cell(row=c_top, column=2 + j, value=f"=Inputs!A{w_first + j}")
    for i in range(k):
        r = c_top + 1 + i
        ws.cell(row=r, column=1, value=f"=Inputs!A{w_first + i}").font = BOLD
        for j in range(k):
            c = ws.cell(row=r, column=2 + j, value=f"=CORREL({rng(stock_cols[i])},{rng(stock_cols[j])})")
            c.font, c.number_format = BLACK, "0.00"
    ws.conditional_formatting.add(
        f"B{c_top + 1}:{L(1 + k)}{c_top + k}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                       end_type="num", end_value=1, end_color="F8696B"))

    # stress scenarios
    x_top = c_top + k + 4
    ws.cell(row=x_top - 1, column=1, value="Market stress scenarios (beta approximation)").font = BOLD
    header(ws, x_top, ["SPY shock", "Est. portfolio move", "Est. loss ($)"])
    for i, shock in enumerate([-0.05, -0.10, -0.20, -0.30]):
        r = x_top + 1 + i
        c = ws[f"A{r}"]
        c.value, c.font, c.fill, c.number_format = shock, BLUE, INPUT_FILL, PCT1
        ws[f"B{r}"], ws[f"B{r}"].number_format, ws[f"B{r}"].font = f"=A{r}*$B$13", PCT, BLACK
        ws[f"C{r}"], ws[f"C{r}"].number_format, ws[f"C{r}"].font = f"=B{r}*{val}", USD, BLACK
    ws.cell(row=x_top + 5, column=1,
            value="Single factor: ignores non-linearity and correlations rising in a crash, so real losses can be worse.").font = \
        Font(name=FONT, italic=True, color="595959")

    # return histogram
    h_top = x_top + 9
    ws.cell(row=h_top - 1, column=1, value="Distribution of daily portfolio returns").font = BOLD
    header(ws, h_top, ["Bin", "Days", "% of days"])
    edges = [round(-0.05 + 0.005 * i, 4) for i in range(21)]     # -5% .. +5%
    bins = [("< -5.0%", f'=COUNTIF({P},"<{edges[0]}")')]
    for lo, hi in zip(edges[:-1], edges[1:]):
        bins.append((f"{lo:+.1%}", f'=COUNTIFS({P},">={lo}",{P},"<{hi}")'))
    bins.append((">= 5.0%", f'=COUNTIF({P},">={edges[-1]}")'))
    for i, (lab, f) in enumerate(bins):
        r = h_top + 1 + i
        ws[f"A{r}"], ws[f"A{r}"].font = lab, BLACK
        ws[f"B{r}"], ws[f"B{r}"].font = f, BLACK
        ws[f"C{r}"], ws[f"C{r}"].font, ws[f"C{r}"].number_format = f"=B{r}/$B$5", BLACK, PCT1
    h_first, h_last = h_top + 1, h_top + len(bins)

    for col, w in zip("ABCDEFGH", [36, 16, 44, 14, 16, 18, 20, 18]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[s_top].height = 32

    # ---------------- Summary ----------------
    ws = ws_sum
    ws["A1"], ws["A1"].font = "Equity risk dashboard", TITLE
    ws["A2"], ws["A2"].font = banner, WARN if synthetic else BLACK
    ws["A3"] = f"10-stock portfolio, last {n} trading days. Change weights on the Inputs tab."
    ws["A3"].font = BLACK
    header(ws, 5, ["Key metric", "Value"])
    keys = [("Annualised return", "B7", PCT), ("Annualised volatility", "B8", PCT), ("Sharpe ratio", "B9", "0.00"),
            ("Maximum drawdown", "B12", PCT), ("Beta vs SPY", "B13", "0.00"),
            ("1-day historical VaR ($)", "B17", USD), ("1-day CVaR ($)", "B18", USD),
            ("Horizon parametric VaR ($)", "B19", USD)]
    for i, (lab, ref, fmt) in enumerate(keys, start=6):
        ws[f"A{i}"], ws[f"A{i}"].font = lab, BLACK
        ws[f"B{i}"] = f"=Risk!{ref}"
        ws[f"B{i}"].font, ws[f"B{i}"].number_format = GREEN, fmt
    ws["A15"], ws["A15"].font = "Weights check", BLACK
    ws["B15"], ws["B15"].font = f"=Inputs!E{w_last + 1}", GREEN
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16

    lc = LineChart()
    lc.title, lc.height, lc.width = "Growth of 1 in the portfolio", 8, 16
    lc.add_data(Reference(ws_ret, min_col=4 + k, min_row=5, max_row=last), titles_from_data=True)
    lc.set_categories(Reference(ws_ret, min_col=1, min_row=first, max_row=last))
    lc.x_axis.number_format, lc.x_axis.majorTimeUnit = "mmm-yy", "months"
    lc.legend = None
    lc.x_axis.delete = lc.y_axis.delete = False
    ws.add_chart(lc, "D5")

    bc = BarChart()
    bc.title, bc.height, bc.width = "Daily return distribution (days per bin, labelled by lower edge)", 8, 16
    bc.add_data(Reference(ws_risk, min_col=2, min_row=h_top, max_row=h_last), titles_from_data=True)
    bc.set_categories(Reference(ws_risk, min_col=1, min_row=h_first, max_row=h_last))
    bc.legend = None
    bc.x_axis.delete = bc.y_axis.delete = False
    ws.add_chart(bc, "D22")

    dc = LineChart()
    dc.title, dc.height, dc.width = "Portfolio drawdown", 8, 16
    dc.add_data(Reference(ws_ret, min_col=6 + k, min_row=5, max_row=last), titles_from_data=True)
    dc.set_categories(Reference(ws_ret, min_col=1, min_row=first, max_row=last))
    dc.x_axis.number_format = "mmm-yy"
    dc.x_axis.tickLblPos = "low"
    dc.legend = None
    dc.x_axis.delete = dc.y_axis.delete = False
    ws.add_chart(dc, "D39")

    # ---------------- Sector_Stats ----------------
    ws = ws_sec
    ws["A1"], ws["A1"].font = "Sector performance (equal-weighted sector portfolios, full history)", TITLE
    ws["A2"], ws["A2"].font = f"Source: sql/queries/q2_sector_performance.sql. {banner}", WARN if synthetic else BLACK
    q2 = pd.read_csv(OUT / "sql" / "q2_sector_performance.csv")
    write_df(ws, q2, 4)
    ws.column_dimensions["A"].width = 26
    for j in range(2, 7):
        ws.column_dimensions[L(j)].width = 18

    # ---------------- Regime_Analysis ----------------
    ws = ws_reg
    ws["A1"], ws["A1"].font = "Sector behaviour by prior-day VIX regime", TITLE
    ws["A2"], ws["A2"].font = f"Source: sql/queries/q6_regime_performance.sql. {banner}", WARN if synthetic else BLACK
    q6 = pd.read_csv(OUT / "sql" / "q6_regime_performance.csv")
    r0 = 4
    for title, col in [("Annualised return (%)", "ann_return_pct"), ("Annualised volatility (%)", "ann_vol_pct")]:
        ws.cell(row=r0, column=1, value=title).font = BOLD
        pv = q6.pivot(index="sector", columns="regime", values=col).reset_index()
        end = write_df(ws, pv, r0 + 1)
        ws.conditional_formatting.add(
            f"B{r0 + 2}:D{end}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                           mid_color="FFFFFF", end_type="max", end_color="63BE7B") if col == "ann_return_pct"
            else ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F8696B"))
        r0 = end + 3
    ws.column_dimensions["A"].width = 26
    for j in range(2, 5):
        ws.column_dimensions[L(j)].width = 18

    # ---------------- Model_Results ----------------
    ws = ws_mod
    ws["A1"], ws["A1"].font = "Volatility forecasting and vol-target backtest (SPY)", TITLE
    ws["A2"], ws["A2"].font = banner, WARN if synthetic else BLACK
    ws["A3"] = "Lower QLIKE / MSE is better. DM stat < 0 means better than GARCH. Walk-forward, refit quarterly."
    ws["A3"].font = BLACK
    mm = pd.read_csv(OUT / "model_metrics.csv")
    end = write_df(ws, mm, 5, fmts={c: "0.0000" for c in mm.columns if c != "model"})
    ws.cell(row=end + 3, column=1, value="QLIKE by prior-day VIX bucket").font = BOLD
    br = pd.read_csv(OUT / "model_metrics_by_regime.csv").pivot(
        index="model", columns="vix_bucket", values="QLIKE").reset_index()
    end = write_df(ws, br, end + 4, fmts={c: "0.0000" for c in br.columns if c != "model"})
    ws.cell(row=end + 3, column=1, value="Volatility-targeting backtest").font = BOLD
    bt = pd.read_csv(OUT / "backtest_metrics.csv")
    pct_cols = ["CAGR", "Ann. vol", "Max drawdown", "Avg SPY weight"]
    write_df(ws, bt, end + 4, fmts={**{c: "0.0%" for c in pct_cols}, "Sharpe": "0.00", "Calmar": "0.00",
                                    "Turnover / yr": "0.0"})
    ws.column_dimensions["A"].width = 28
    for j in range(2, 9):
        ws.column_dimensions[L(j)].width = 16

    OUT.mkdir(exist_ok=True)
    path = OUT / "risk_dashboard.xlsx"
    wb.save(path)
    print(f"saved {path}")


if __name__ == "__main__":
    main()
